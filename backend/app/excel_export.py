"""Builds the KF & Futura master-format Excel workbook."""

from __future__ import annotations

import io
from datetime import datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .list_formatter import format_single_academic_line, sort_academic_qualifications

MASTER_HEADERS = [
    "NAME", "Company Name (Futura and KF Bioplant)", "BIRTH_DT", "Gender", "Marital Status", "DESIG",
    "Designation Code", "F_Designation Code", "JOIN_DT", "PRB_DT", "CONF_DT",
    "DEPT_NAME", "Department Code", "F_Department Code", "Lab Details", "Section Code", "F_Section Code",
    "EPAN_NO", "EMAILID", "Address 1", "Address2",
    "Pin 1", "Pin 2", "ADHAR_NO", "PF_UAN", "ESI Number", "EMP_Mobile Number", "Alternate_Mobile Number",
    "Father Name", "Father DOB", "Mother Name", "Mother DOB", "Spouse_Name", "Spouse DOB",
    "Child_1", "Child 1 DOB", "Child_2", "Child 2 DOB",
    "Qualification", "Previous Employement", "Present Employement", "Total Experience", "Reference Name",
]

# Multi-line bulleted list columns get extra width
WIDE_LIST_COLUMNS = {
    MASTER_HEADERS.index("Qualification"),
    MASTER_HEADERS.index("Previous Employement"),
    MASTER_HEADERS.index("Present Employement"),
}

HISTORY_HEADERS = [
    "Employee ID", "Employee Name", "Company", "Designation", "Designation Code", "Department",
    "Department Code", "Section", "Section Code", "Highest Qualification", "Academics Breakdown",
    "Total Experience", "Salary / CTC", "Previous Company", "Languages Known",
]


def _fill(color: str) -> PatternFill:
    return PatternFill(fill_type="solid", start_color=color, end_color=color)


def _border(color: str, bottom_style: str = "thin") -> Border:
    side = Side(style="thin", color=color)
    return Border(top=side, left=side, right=side, bottom=Side(style=bottom_style, color=color))


GRID_BORDER = _border("E0E0E0")


def _v(rec: dict, key: str, default: str = "") -> Any:
    value = rec.get(key)
    return default if value in (None, "") else value


def _master_row(rec: dict) -> list[Any]:
    return [
        _v(rec, "NAME"), _v(rec, "COMPANY"), _v(rec, "BIRTH_DT"), _v(rec, "Gender"), _v(rec, "MaritalStatus"),
        _v(rec, "DESIG"), _v(rec, "DESIG_CODE", "-"), _v(rec, "F_DESIG_CODE", "-"),
        _v(rec, "JOIN_DT"), _v(rec, "PRB_DT", "-"), _v(rec, "CONF_DT", "-"), _v(rec, "DEPT_NAME"),
        _v(rec, "DEPT_CODE", "-"), _v(rec, "F_DEPT_CODE", "-"), _v(rec, "LabDetails"),
        _v(rec, "SECTION_CODE", "-"), _v(rec, "F_SECTION_CODE", "-"), _v(rec, "EPAN_NO", "-"), _v(rec, "EMAILID"),
        _v(rec, "Address1"), _v(rec, "Address2"), _v(rec, "Pin1"),
        _v(rec, "Pin 2") or _v(rec, "Pin2"), _v(rec, "ADHAR_NO", "-"), _v(rec, "PF_UAN", "-"),
        _v(rec, "ESI_Number", "-"), _v(rec, "EMP_MobileNumber"),
        _v(rec, "Alternate_MobileNumber", "-"), _v(rec, "FatherName"), _v(rec, "FatherDOB"), _v(rec, "MotherName"),
        _v(rec, "MotherDOB"),
        _v(rec, "SpouseName", "-"), _v(rec, "SpouseDOB", "-"), _v(rec, "Child1", "-"), _v(rec, "Child1DOB", "-"),
        _v(rec, "Child2", "-"), _v(rec, "Child2DOB", "-"),
        _v(rec, "Qualification"), _v(rec, "PreviousEmployment"), _v(rec, "PresentEmployment"),
        _v(rec, "TotalExperience", "-"),
        _v(rec, "ReferenceName"),
    ]


def build_employee_workbook(records: list[dict]) -> bytes:
    wb = Workbook()
    wb.properties.creator = "KF & Futura Employee Data Portal"
    wb.properties.lastModifiedBy = "Employee Form AI Scanner"

    # ---------------- Sheet 1: Master Employee Data ----------------
    ws = wb.active
    ws.title = "Master Employee Data"
    ws.sheet_view.showGridLines = True
    last_col = get_column_letter(len(MASTER_HEADERS))

    ws.merge_cells(f"A1:{last_col}1")
    title = ws["A1"]
    title.value = "KF & FUTURA - EMPLOYEE PERSONAL DATA MASTER RECORD"
    title.font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    title.fill = _fill("1B5E20")
    title.alignment = Alignment(vertical="center", horizontal="center")
    ws.row_dimensions[1].height = 32

    ws.merge_cells(f"A2:{last_col}2")
    sub = ws["A2"]
    sub.value = f"Generated on: {datetime.now().strftime('%d/%m/%Y, %I:%M:%S %p')} | Total Employees: {len(records)}"
    sub.font = Font(name="Calibri", size=10, italic=True, color="333333")
    sub.fill = _fill("E8F5E9")
    sub.alignment = Alignment(vertical="center", horizontal="left", indent=1)
    ws.row_dimensions[2].height = 20

    header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    header_fill = _fill("2E7D32")
    header_border = _border("1B5E20", bottom_style="medium")
    for col_idx, header in enumerate(MASTER_HEADERS, start=1):
        cell = ws.cell(row=3, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
        cell.border = header_border
    ws.row_dimensions[3].height = 30

    body_font = Font(name="Calibri", size=10)
    body_alignment = Alignment(vertical="center", horizontal="left", wrap_text=True)
    max_lens = [len(h) for h in MASTER_HEADERS]

    for idx, rec in enumerate(records):
        values = _master_row(rec)
        row_num = 4 + idx
        fill = _fill("FFFFFF" if idx % 2 == 0 else "F1F8E9")
        max_lines = 1
        for col_idx, value in enumerate(values):
            text = str(value)
            cell = ws.cell(row=row_num, column=col_idx + 1, value=value)
            cell.font = body_font
            cell.fill = fill
            cell.alignment = body_alignment
            cell.border = GRID_BORDER
            lines = text.split("\n")
            max_lines = max(max_lines, len(lines))
            max_lens[col_idx] = max(max_lens[col_idx], *(len(line) for line in lines))
        ws.row_dimensions[row_num].height = max(26, max_lines * 18)

    for col_idx, max_len in enumerate(max_lens):
        if col_idx in WIDE_LIST_COLUMNS:
            width = max(min(max_len + 4, 60), 45)
        else:
            width = max(min(max_len + 4, 45), 12)
        ws.column_dimensions[get_column_letter(col_idx + 1)].width = width

    # ---------------- Sheet 2: Academic & Detailed History ----------------
    ws2 = wb.create_sheet("Academic & Detailed History")
    ws2.append(HISTORY_HEADERS)
    for cell in ws2[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _fill("1B5E20")
    ws2.row_dimensions[1].height = 28

    for rec in records:
        ext = rec.get("extendedDetails") if isinstance(rec.get("extendedDetails"), dict) else {}
        academics = sort_academic_qualifications(ext.get("academics") or [])
        academic_str = "\n".join(format_single_academic_line(a) for a in academics if isinstance(a, dict))
        languages = ", ".join(
            str(lang.get("language")) for lang in (ext.get("languages") or []) if isinstance(lang, dict) and lang.get("language")
        )
        ws2.append([
            rec.get("id"), rec.get("NAME"), rec.get("COMPANY") or "-", rec.get("DESIG"), rec.get("DESIG_CODE"),
            rec.get("DEPT_NAME"), rec.get("DEPT_CODE"), rec.get("LabDetails"), rec.get("SECTION_CODE"),
            rec.get("Qualification"), academic_str or rec.get("Qualification"),
            rec.get("TotalExperience") or "-",
            ext.get("ctcGrossPerAnnum") or "N/A", rec.get("PreviousEmployment"), languages or "N/A",
        ])
        row_num = ws2.max_row
        previous = str(rec.get("PreviousEmployment") or "")
        lines = max(1, len(academic_str.split("\n")) if academic_str else 1, len(previous.split("\n")) if previous else 1)
        ws2.row_dimensions[row_num].height = max(24, lines * 18)
        for cell in ws2[row_num]:
            if cell.value is not None:
                cell.alignment = body_alignment
                cell.border = GRID_BORDER

    for col_idx in range(1, len(HISTORY_HEADERS) + 1):
        ws2.column_dimensions[get_column_letter(col_idx)].width = 30

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
